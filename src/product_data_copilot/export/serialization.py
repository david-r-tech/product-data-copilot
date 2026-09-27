"""Spreadsheet output that preserves text instead of executing it as formulas."""

import csv
import json
import math
import re
from io import BytesIO, StringIO
from numbers import Number

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

INVALID_XML_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
FORMULA_PREFIXES = ("=", "+", "-", "@", "＝", "＋", "－", "＠")


class ExportError(ValueError):
    """A workbook limit or serialization problem the UI can explain safely."""


def _cell_value(value):
    if isinstance(value, (list, tuple, dict, set)):
        value = json.dumps(list(value) if isinstance(value, set) else value, ensure_ascii=False, default=str)
    if value is None or pd.isna(value):
        return ""
    if isinstance(value, Number) and not isinstance(value, bool):
        return value if math.isfinite(value) else str(value)
    text = str(value)
    return INVALID_XML_CHARS.sub(lambda match: f"\\u{ord(match.group()):04x}", text)


def safe_csv_bytes(dataframe):
    """Return a UTF-8 BOM CSV with formula-like text escaped for spreadsheets."""
    def escape(value):
        value = _cell_value(value)
        if isinstance(value, str) and (
            value.lstrip().startswith(FORMULA_PREFIXES) or value.startswith(("\t", "\r", "\n"))
        ):
            return "'" + value
        return value

    output = StringIO(newline="")
    writer = csv.writer(output, quoting=csv.QUOTE_ALL)
    writer.writerow([escape(column) for column in dataframe.columns])
    for row in dataframe.itertuples(index=False, name=None):
        writer.writerow([escape(value) for value in row])
    return output.getvalue().encode("utf-8-sig")


def workbook_bytes(sheets):
    """Serialize derived sheets without mutating inputs or silently truncating text."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="183B56")
    for sheet_name, dataframe in sheets.items():
        sheet = workbook.create_sheet(sheet_name)
        sheet.freeze_panes = "A2"
        rows = [list(dataframe.columns)]
        rows.extend(dataframe.itertuples(index=False, name=None))
        widths = {}
        for row_index, row in enumerate(rows, start=1):
            for column_index, value in enumerate(row, start=1):
                value = _cell_value(value)
                if isinstance(value, str) and len(value) > 32767:
                    raise ExportError(
                        f"A cell in {sheet_name} exceeds Excel's 32,767-character limit. "
                        "Shorten the source text before exporting."
                    )
                cell = sheet.cell(row_index, column_index, value)
                if isinstance(value, str):
                    cell.data_type = "s"
                    cell.number_format = "@"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_index == 1:
                    cell.fill = header_fill
                    cell.font = Font(color="FFFFFF", bold=True)
                widths[column_index] = max(widths.get(column_index, 12), min(55, len(str(value)) + 2))
        if len(dataframe.columns):
            sheet.auto_filter.ref = sheet.dimensions
        for column_index, width in widths.items():
            sheet.column_dimensions[get_column_letter(column_index)].width = width
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
