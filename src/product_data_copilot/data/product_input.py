"""Read bounded local product files without coercing identifiers to numbers."""

import csv
import hashlib
from io import BytesIO, StringIO
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_PRODUCTS = 1000
MAX_COLUMNS = 100
REQUIRED_COLUMNS = {"sku", "product_name"}


class ProductInputError(ValueError):
    """An input problem that can be shown directly to the user."""


def dataset_fingerprint(content, source_name):
    """Scope review state to the complete input, including its source identity."""
    return hashlib.sha256(source_name.encode("utf-8") + b"\0" + content).hexdigest()


def validate_product_identity(products):
    """Reject ambiguous product identity before issues or decisions are assigned."""
    if "sku" not in products:
        raise ProductInputError("Die Datei muss eine Spalte namens sku enthalten.")
    skus = products["sku"].fillna("").astype(str).str.strip()
    if skus.eq("").any():
        rows = [str(i + 2) for i, blank in enumerate(skus.eq("")) if blank]
        raise ProductInputError(
            "Jeder Artikel benötigt eine SKU. Sie fehlt in den Datenzeilen: "
            + ", ".join(rows[:10]) + ". Korrigiere die Datei und lade sie erneut hoch."
        )
    if skus.duplicated(keep=False).any():
        raise ProductInputError(
            "SKUs müssen eindeutig sein. Es wurden doppelte SKUs gefunden; korrigiere "
            "die Datei vor der Produktprüfung."
        )


def load_product_file(content, filename):
    """Load UTF-8 CSV or the first XLSX sheet; formulas remain literal text.

    All imported cells are text. Column names and SKUs are trimmed. Other cell
    text is retained; optional absent columns are handled by the rule engine.
    """
    if not content:
        raise ProductInputError("Die Datei ist leer. Lade eine Datei mit Produktzeilen hoch.")
    if len(content) > MAX_FILE_BYTES:
        raise ProductInputError("Die Datei überschreitet die Grenze dieses lokalen MVP von 10 MB.")
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        rows = _read_csv_rows(content)
    elif suffix == ".xlsx":
        rows = _read_excel_rows(content)
    else:
        raise ProductInputError("Verwende eine CSV- oder XLSX-Datei.")
    if not rows:
        raise ProductInputError("Die Datei enthält keine Kopfzeile oder Produktzeilen.")
    headers = [str(value).strip().lower() for value in rows[0]]
    if len(headers) > MAX_COLUMNS:
        raise ProductInputError("Verwende höchstens 100 Spalten pro Produktdatei.")
    if any(not header for header in headers) or len(set(headers)) != len(headers):
        raise ProductInputError("Spaltennamen müssen ausgefüllt und eindeutig sein.")
    missing = sorted(REQUIRED_COLUMNS - set(headers))
    if missing:
        raise ProductInputError(
            "Erforderliche Spalte(n) fehlen: " + ", ".join(missing)
            + ". Verwende die Beispieldatei als Spaltenvorlage."
        )
    records = []
    for row_number, row in enumerate(rows[1:], start=2):
        if not any(str(value).strip() for value in row):
            continue
        if len(row) != len(headers):
            raise ProductInputError(
                f"Zeile {row_number} hat {len(row)} Zellen; erwartet werden {len(headers)}. "
                "Prüfe das Trennzeichen und setze Texte mit Trennzeichen in Anführungszeichen."
            )
        records.append(row)
    if not records:
        raise ProductInputError("Die Datei enthält Spaltenüberschriften, aber keine Produktzeilen.")
    if len(records) > MAX_PRODUCTS:
        raise ProductInputError("Verwende in diesem lokalen MVP höchstens 1.000 Artikel pro Datei.")
    products = pd.DataFrame(records, columns=headers, dtype=str)
    products["sku"] = products["sku"].str.strip()
    validate_product_identity(products)
    return products


def _read_csv_rows(content):
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ProductInputError("Speichere die CSV als UTF-8 und lade sie erneut hoch.") from error
    header_line = text.splitlines()[0] if text.splitlines() else ""
    try:
        delimiter = csv.Sniffer().sniff(header_line, delimiters=",;\t").delimiter
    except csv.Error:
        delimiter = ","
    try:
        return list(csv.reader(StringIO(text, newline=""), delimiter=delimiter, strict=True))
    except csv.Error as error:
        raise ProductInputError("Die CSV ist fehlerhaft. Prüfe Anführungszeichen und Trennzeichen.") from error


def _read_excel_rows(content):
    workbook = None
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise ProductInputError("Die entpackte Arbeitsmappe ist für dieses lokale MVP zu groß.")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
        sheet = workbook.worksheets[0]
        # Do not trust inflated worksheet dimensions written by other tools.
        sheet.reset_dimensions()
        rows = []
        for row in sheet.iter_rows(values_only=True):
            if len(row) > MAX_COLUMNS:
                raise ProductInputError("Verwende höchstens 100 Spalten pro Produktdatei.")
            rows.append(["" if value is None else str(value) for value in row])
            if len(rows) > MAX_PRODUCTS + 1:
                raise ProductInputError("Verwende in diesem lokalen MVP höchstens 1.000 Artikel pro Datei.")
        # XLSX rows omit trailing empty cells. Pad to the header width only.
        if rows:
            rows = [row + [""] * max(0, len(rows[0]) - len(row)) for row in rows]
        return rows
    except ProductInputError:
        raise
    except (BadZipFile, InvalidFileException, ParseError, ValueError, KeyError, IndexError, OSError) as error:
        raise ProductInputError(
            "Die Excel-Datei konnte nicht gelesen werden. Lade eine gültige, unverschlüsselte XLSX-Arbeitsmappe hoch."
        ) from error
    finally:
        if workbook is not None:
            workbook.close()
