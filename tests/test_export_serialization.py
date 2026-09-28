import csv
from io import BytesIO, StringIO

import pandas as pd
import pytest
from openpyxl import load_workbook

import app
from product_data_copilot.export.serialization import ExportError, safe_csv_bytes, workbook_bytes


def test_excel_keeps_formula_like_strings_identifiers_and_errors_as_text():
    data = pd.DataFrame({"sku": ["000123"] * 6, "text": ["=1+1", "+1+1", "-2+3", "@SUM(A1)", "#N/A", "Größe"]})
    original = data.copy(deep=True)
    workbook = load_workbook(BytesIO(workbook_bytes({"Data": data})), data_only=False)
    for index, value in enumerate(data["text"], start=2):
        assert workbook["Data"].cell(index, 2).value == value
        assert workbook["Data"].cell(index, 2).data_type == "s"
        assert workbook["Data"].cell(index, 1).value == "000123"
    pd.testing.assert_frame_equal(data, original)


def test_controls_are_visible_escapes_and_long_text_is_not_silently_truncated():
    data = pd.DataFrame({"description": ["Before\x0bAfter"]})
    workbook = load_workbook(BytesIO(workbook_bytes({"Data": data})))
    assert workbook["Data"]["A2"].value == r"Before\u000bAfter"
    assert data.iloc[0, 0] == "Before\x0bAfter"
    with pytest.raises(ExportError, match="32,767"):
        workbook_bytes({"Data": pd.DataFrame({"text": ["x" * 32768]})})


@pytest.mark.parametrize("text", ["=1+1", " +1", "-1+2", "@SUM(A1)", "\ttext", "\n=1", "＝1+1"])
def test_csv_escapes_formula_text_and_retains_unicode(text):
    data = pd.DataFrame({"description": [text], "name": ["Größe"]})
    content = safe_csv_bytes(data)
    assert content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(StringIO(content.decode("utf-8-sig"))))
    assert rows[1] == ["'" + text, "Größe"]
    assert data.iloc[0, 0] == text


def test_management_workbook_preserves_source_data(valid_product):
    valid_product["description"] = "=1+1"
    products = pd.DataFrame([valid_product])
    before = products.copy(deep=True)
    issues = app.find_product_issues(products)
    scores = app.calculate_readiness_scores(products, issues)
    tasks = app.create_review_tasks(issues, scores)
    management = load_workbook(BytesIO(app.create_excel_management_export(
        "Test input", products, scores, issues, tasks,
    )))
    assert management.sheetnames == app.MANAGEMENT_EXPORT_SHEETS
    assert management["Source Products"]["D2"].value == "=1+1"
    assert management["Source Products"]["D2"].data_type == "s"
    assert "AI Suggestions" not in management.sheetnames
    pd.testing.assert_frame_equal(products, before)
