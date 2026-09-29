"""Focused regression tests for management-report construction."""

import inspect
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

import app
from product_data_copilot.export import management_report
from product_data_copilot.export.export_helpers import MANAGEMENT_EXPORT_SHEETS


def _report_inputs():
    products = pd.DataFrame([
        {"sku": "A", "product_name": "First", "description": "=1+1"},
        {"sku": "B", "product_name": "Second", "description": "Text"},
        {"sku": "C", "product_name": "Third", "description": "Text"},
    ])
    issues = pd.DataFrame([
        {"sku": "A", "severity": "Critical", "issue_type": "Missing ean"},
        {"sku": "A", "severity": "Warning", "issue_type": "Missing brand"},
        {"sku": "B", "severity": "Info", "issue_type": "Missing image_url"},
        {"sku": "B", "severity": "Warning", "issue_type": "Short description"},
    ])
    scores = pd.DataFrame([
        {"sku": "A", "overall_readiness_score": 90, "readiness_status": "Ready"},
        {"sku": "B", "overall_readiness_score": 75, "readiness_status": "Needs Review"},
        {"sku": "C", "overall_readiness_score": 50, "readiness_status": "Critical"},
    ])
    tasks = pd.DataFrame([
        {"sku": "A", "priority": "High"},
        {"sku": "A", "priority": "Medium"},
        {"sku": "B", "priority": "Low"},
        {"sku": "B", "priority": "Medium"},
    ])
    return products, issues, scores, tasks


def test_app_uses_extracted_management_report_pipeline():
    assert app.create_management_summary is management_report.create_management_summary
    assert app.create_excel_management_export is management_report.create_excel_management_export
    app_source = inspect.getsource(app)
    assert "def create_management_summary" not in app_source
    assert "def create_excel_management_export" not in app_source


def test_management_summary_values_remain_unchanged():
    products, issues, scores, tasks = _report_inputs()
    summary = management_report.create_management_summary(
        "Test input", products, issues, scores, tasks,
    )

    assert summary.to_records(index=False).tolist() == [
        ("Data source", "Test input"),
        ("Total products", 3),
        ("Total issues", 4),
        ("Critical issues", 1),
        ("Warning issues", 2),
        ("Info issues", 1),
        ("Products affected by issues", 2),
        ("Average overall readiness score", 72),
        ("Ready products", 1),
        ("Needs Review products", 1),
        ("Critical products", 1),
        ("Total review tasks", 4),
        ("High priority tasks", 1),
        ("Medium priority tasks", 2),
        ("Low priority tasks", 1),
    ]


def test_management_workbook_sheets_and_content_remain_unchanged():
    products, issues, scores, tasks = _report_inputs()
    sheets = management_report.prepare_management_workbook_sheets(
        "Test input", products, scores, issues, tasks,
    )

    assert list(sheets) == MANAGEMENT_EXPORT_SHEETS
    assert sheets["Product Scores"] is scores
    assert sheets["Issues"] is issues
    assert sheets["Review Tasks"] is tasks
    assert sheets["Source Products"] is products

    workbook = load_workbook(BytesIO(management_report.create_excel_management_export(
        "Test input", products, scores, issues, tasks,
    )))
    assert workbook.sheetnames == MANAGEMENT_EXPORT_SHEETS
    assert workbook["Management Summary"]["A2"].value == "Data source"
    assert workbook["Management Summary"]["B2"].value == "Test input"
    assert workbook["Source Products"]["C2"].value == "=1+1"
    assert workbook["Source Products"]["C2"].data_type == "s"


def test_management_export_delegates_to_existing_safe_serialization(monkeypatch):
    products, issues, scores, tasks = _report_inputs()
    calls = []

    def fake_workbook_bytes(sheets):
        calls.append(sheets)
        return b"serialized"

    monkeypatch.setattr(management_report, "workbook_bytes", fake_workbook_bytes)
    result = management_report.create_excel_management_export(
        "Test input", products, scores, issues, tasks,
    )

    assert result == b"serialized"
    assert len(calls) == 1
    assert list(calls[0]) == MANAGEMENT_EXPORT_SHEETS
